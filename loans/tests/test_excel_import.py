"""Excel import aligned with current Region/Zone/City, user, product, and loan models."""

from __future__ import annotations

import tempfile
from decimal import Decimal
from pathlib import Path

import pandas as pd
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse

from collateral.models import MainWork, SubWorkUnitPrice
from loans.excel_import import import_pack, run_import
from loans.models import (
    ApprovalCommitteeLevel,
    Branch,
    City,
    CollateralType,
    CustomUser,
    District,
    LoanApplicationDocumentType,
    LoanCategory,
    LoanRequest,
    Region,
    Zone,
)


def _xlsx(rows, path, sheet='Sheet1'):
    pd.DataFrame(rows).to_excel(path, index=False, sheet_name=sheet)


class ExcelImportTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_import_zones_requires_region_not_district(self):
        path = self.dir / 'zones.xlsx'
        _xlsx([{'name': 'Mekelle District'}], path)
        result = run_import('zones', path)
        self.assertTrue(result.errors)
        self.assertEqual(Zone.objects.count(), 0)
        self.assertEqual(District.objects.count(), 0)

    def test_geography_and_branch_chain(self):
        run_import('regions', self._file('r.xlsx', [{'name': 'Tigray'}]))
        run_import('zones', self._file('z.xlsx', [{'region': 'Tigray', 'name': 'Mekelle Zone'}]))
        run_import('cities', self._file('c.xlsx', [
            {'region': 'Tigray', 'zone': 'Mekelle Zone', 'name': 'Mekelle'},
        ]))
        run_import('districts', self._file('d.xlsx', [{'name': 'Mekelle District'}]))
        # legacy “zone” column means operational district
        run_import('branches', self._file('b.xlsx', [
            {'zone': 'Mekelle District', 'name': 'Mekelle Main'},
        ]))
        self.assertTrue(Region.objects.filter(name='Tigray').exists())
        self.assertTrue(Zone.objects.filter(name='Mekelle Zone', region__name='Tigray').exists())
        self.assertTrue(City.objects.filter(name='Mekelle').exists())
        branch = Branch.objects.get(name='Mekelle Main')
        self.assertEqual(branch.district.name, 'Mekelle District')

    def test_categories_collateral_and_users(self):
        District.objects.create(name='Mekelle District')
        Branch.objects.create(name='Mekelle Main', district=District.objects.get(name='Mekelle District'))
        run_import('loan_categories', self._file('cat.xlsx', [
            {'name': 'MSME Trade', 'appraisal_mode': 'msme'},
            {'name': 'Corporate CapEx', 'appraisal_mode': 'corporate', 'product_family': 'project'},
        ]))
        run_import('collateral_types', self._file('col.xlsx', [
            {'collateral': 'Building / House', 'kind': 'building'},
            {'name': 'Land', 'kind': 'land'},
        ]))
        result = run_import('users', self._file('u.xlsx', [
            {
                'username': 'bm.mekele',
                'email': 'bm@decsi.test',
                'phone_number': '0914000001',
                'role': 'Branch Manager',
                'district': 'Mekelle District',
                'branch': 'Mekelle Main',
            },
            {
                'username': 'credit.head',
                'email': 'ch@decsi.test',
                'phone_number': '0914000002',
                'role': 'credit_head',
                'department': 'credit',
            },
        ]), default_password='Temp@12345')
        self.assertFalse(result.errors, result.errors)
        bm = CustomUser.objects.get(username='bm.mekele')
        self.assertEqual(bm.role, 'branch_manager')
        self.assertTrue(bm.check_password('Temp@12345'))
        head = CustomUser.objects.get(username='credit.head')
        self.assertIsNone(head.branch_id)
        self.assertEqual(head.department.key, 'credit')
        self.assertEqual(LoanCategory.objects.get(name='Corporate CapEx').appraisal_mode, 'corporate')
        self.assertEqual(LoanCategory.objects.get(name='Corporate CapEx').product_family, 'project')
        self.assertEqual(LoanCategory.objects.get(name='MSME Trade').product_family, 'general')
        self.assertEqual(CollateralType.objects.get(name='Building / House').kind, 'building')

    def test_loan_request_uses_existing_id_and_customer_number(self):
        district = District.objects.create(name='Mekelle District')
        Branch.objects.create(name='Mekelle Main', district=district)
        LoanCategory.objects.create(name='MSME Trade', appraisal_mode='msme')
        CollateralType.objects.create(name='Land', kind='land')
        CustomUser.objects.create_user(
            username='lo.mekele1', password='x', phone_number='0914000099',
            role='loan_officer', district=district, branch=Branch.objects.get(name='Mekelle Main'),
        )
        result = run_import('loan_requests', self._file('lr.xlsx', [{
            'loan_request_id': 'DECSI-1001',
            'applicant_name': 'Hagos Tesfay',
            'phone_number': '0914111222',
            'customer_number': 'CIF001',
            'category': 'MSME Trade',
            'collateral': 'Land',
            'amount_requested': 250000,
            'reason': 'Working capital',
            'status': 'approved',
            'branch': 'Mekelle Main',
            'customer_history': 'existing',
            'assigned_loan_officer': 'lo.mekele1',
            'origin_level': 'branch',
        }]))
        self.assertFalse(result.errors, result.errors)
        loan = LoanRequest.objects.get(loan_request_id='DECSI-1001')
        self.assertEqual(loan.customer_number, 'CIF001')
        self.assertEqual(loan.status, 'Approved')
        self.assertTrue(loan.queue_approved)
        self.assertEqual(loan.assigned_loan_officer.username, 'lo.mekele1')
        self.assertEqual(loan.amount_requested, Decimal('250000'))

    def test_construction_catalog_and_committee(self):
        region = Region.objects.create(name='Tigray')
        zone = Zone.objects.create(region=region, name='Mekelle Zone')
        City.objects.create(zone=zone, name='Mekelle')
        run_import('construction_catalog', self._file('catl.xlsx', [{
            'main_work': 'Foundation',
            'main_work_order': 1,
            'sub_work': 'Excavation',
            'sub_work_order': '1.1',
            'sub_sub_work': 'Manual excavation',
            'unit_measure': 'm³',
            'city': 'Mekelle',
            'zone': 'Mekelle Zone',
            'unit_price': 1500,
        }]))
        self.assertTrue(MainWork.objects.filter(name='Foundation').exists())
        self.assertEqual(SubWorkUnitPrice.objects.get().unit_price, Decimal('1500'))

        run_import('committee_levels', self._file('cl.xlsx', [{
            'key': 'branch',
            'name': 'Branch Committee',
            'voter_scope': 'branch',
            'sequence_order': 1,
            'min_approvals_required': 2,
            'max_loan_amount': 500000,
        }]))
        run_import('committee_members', self._file('cm.xlsx', [{
            'level_key': 'branch',
            'participant_type': 'role',
            'role': 'branch_manager',
        }]))
        level = ApprovalCommitteeLevel.objects.get(key='branch')
        self.assertTrue(
            level.member_rules.filter(
                participant_type='role', role='branch_manager', is_active=True,
            ).exists()
        )

    def test_document_pack(self):
        LoanCategory.objects.create(name='MSME Trade', appraisal_mode='msme')
        run_import('document_types', self._file('dt.xlsx', [{
            'name': 'National ID',
            'order': 10,
            'is_required': 'TRUE',
            'enable_ocr_match': 'yes',
        }]))
        result = run_import('category_documents', self._file('cd.xlsx', [{
            'category': 'MSME Trade',
            'document_type': 'National ID',
            'is_required': 'TRUE',
            'order': 10,
        }]))
        self.assertFalse(result.errors, result.errors)
        self.assertTrue(
            LoanApplicationDocumentType.objects.get(name='National ID').enable_ocr_match
        )
        self.assertEqual(
            LoanCategory.objects.get(name='MSME Trade').document_requirements.count(), 1,
        )

    def test_update_and_dry_run_command(self):
        path = self._file('reg.xlsx', [{'name': 'Tigray'}])
        call_command('import_regions', str(path))
        self.assertEqual(Region.objects.count(), 1)
        with self.assertRaises(CommandError):
            call_command('import_zones', str(self._file('bad.xlsx', [{'name': 'OnlyName'}])))
        call_command('import_regions', str(path), dry_run=True, update=True)
        self.assertEqual(Region.objects.filter(name='Tigray').count(), 1)

    def test_migration_pack_import(self):
        from openpyxl import Workbook

        path = self.dir / 'pack.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '01_Regions'
        ws.append(['name'])
        ws.append(['Tigray'])
        ws2 = wb.create_sheet('04_Districts')
        ws2.append(['name'])
        ws2.append(['Eastern District'])
        ws3 = wb.create_sheet('Lookups')
        ws3.append(['ignored'])
        wb.save(path)
        result = import_pack(path)
        self.assertFalse(result.errors, result.errors)
        self.assertTrue(Region.objects.filter(name='Tigray').exists())
        self.assertTrue(District.objects.filter(name='Eastern District').exists())

    def test_upload_zones_requires_superuser(self):
        user = CustomUser.objects.create_user(
            username='lo1', password='pass', phone_number='0914000000', role='loan_officer',
        )
        self.client.force_login(user)
        resp = self.client.get(reverse('upload_zones'))
        self.assertIn(resp.status_code, (302, 403))

    def _file(self, name, rows):
        path = self.dir / name
        _xlsx(rows, path)
        return path
