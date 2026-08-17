"""Reporting / exports: pipeline Excel, appraisal pack Excel, branch dashboards."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from django.db.models import Count, Q, QuerySet, Sum
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


REPORT_ROLES = (
    'branch_manager', 'cooperative_manager', 'operation_manager', 'finance_manager',
    'credit_head', 'credit_loan_officer',
    'loan_officer', 'district_manager', 'accountant', 'admin', 'superadmin',
    'ceo', 'vp', 'vp_operations', 'vp_it', 'vp_customer_service',
    'board_member', 'risk_compliance', 'auditor',
    'engineer', 'engineering_head',
)

# Organization-wide MIS (not limited to one branch/district).
# Accountant/auditor are org-wide only when they have no branch/district attachment.
ORG_WIDE_REPORT_ROLES = (
    'admin', 'superadmin', 'cooperative_manager', 'operation_manager', 'finance_manager',
    'credit_head',
    'ceo', 'vp', 'vp_operations', 'vp_it', 'vp_customer_service',
    'board_member', 'risk_compliance',
)


def user_can_access_reports(user) -> bool:
    if not user or not user.is_authenticated:
        return False
    if getattr(user, 'is_superuser', False):
        return True
    return getattr(user, 'role', None) in REPORT_ROLES


def is_org_wide_reporter(user) -> bool:
    if not user:
        return False
    if getattr(user, 'is_superuser', False):
        return True
    role = getattr(user, 'role', None)
    if role in ORG_WIDE_REPORT_ROLES:
        return True
    # Unscoped HO accountant / auditor (no branch or district) sees org-wide.
    if role in ('accountant', 'auditor'):
        if not getattr(user, 'branch_id', None) and not getattr(user, 'district_id', None):
            return True
    return False


def is_engineering_reporter(user) -> bool:
    return getattr(user, 'role', None) in ('engineer', 'engineering_head')


def engineering_queue_q():
    """Loans in the engineering / collateral QA pipeline (shared with collateral views)."""
    from loans.models import LoanRequest

    return (
        Q(sent_to_engineering_at__isnull=False)
        | Q(collateral_submitted_at__isnull=False)
        | Q(collateral_engineering_status__in=(
            LoanRequest.ENG_COLLATERAL_PENDING,
            LoanRequest.ENG_COLLATERAL_RETURNED,
            LoanRequest.ENG_COLLATERAL_APPROVED,
        ))
        | Q(assigned_engineer__isnull=False)
    )


def report_scope_kind(user) -> str:
    """branch | district | assigned | engineering_assigned | engineering_queue | organization | none"""
    if not user or not getattr(user, 'is_authenticated', False):
        return 'none'
    if is_org_wide_reporter(user):
        return 'organization'
    role = getattr(user, 'role', None)
    if role in ('loan_officer', 'credit_loan_officer'):
        return 'assigned'
    if role == 'engineer':
        return 'engineering_assigned'
    if role == 'engineering_head':
        return 'engineering_queue'
    if role in ('accountant', 'auditor', 'branch_manager'):
        if getattr(user, 'branch_id', None):
            return 'branch'
        if getattr(user, 'district_id', None):
            return 'district'
        return 'none'
    if role == 'district_manager':
        return 'district' if getattr(user, 'district_id', None) else 'none'
    return 'none'


def report_scope_label(user) -> str:
    """Human-readable scope for Reports / CI badges."""
    kind = report_scope_kind(user)
    if kind == 'organization':
        return 'Organization-wide'
    if kind == 'assigned':
        return 'Assigned to you'
    if kind == 'engineering_assigned':
        return 'Assigned to you (engineering)'
    if kind == 'engineering_queue':
        return 'Engineering / collateral queue'
    if kind == 'branch':
        branch = getattr(user, 'branch', None)
        return branch.name if branch else 'Your branch'
    if kind == 'district':
        district = getattr(user, 'district', None)
        if district:
            return f'District: {district.name}'
        return 'Your district'
    return 'No scope assigned'


def reporting_base_queryset(user) -> QuerySet:
    from loans.models import LoanRequest

    qs = LoanRequest.objects.select_related(
        'branch', 'branch__district', 'category', 'collateral',
        'assigned_loan_officer', 'assigned_engineer', 'appraisal',
    ).order_by('-date_requested', '-id')
    role = getattr(user, 'role', None)
    if is_org_wide_reporter(user):
        return qs
    if role in ('loan_officer', 'credit_loan_officer'):
        return qs.filter(assigned_loan_officer=user)
    if role == 'engineer':
        return qs.filter(assigned_engineer=user)
    if role == 'engineering_head':
        return qs.filter(engineering_queue_q())
    if role in ('branch_manager', 'accountant', 'auditor'):
        if getattr(user, 'branch_id', None):
            return qs.filter(branch_id=user.branch_id)
        if getattr(user, 'district_id', None):
            return qs.filter(branch__district_id=user.district_id)
        return qs.none()
    if role == 'district_manager' and getattr(user, 'district_id', None):
        return qs.filter(branch__district_id=user.district_id)
    return qs.none()


def _branch_in_user_scope(user, branch_id: str) -> bool:
    """True if branch_id is within the user's reporting scope."""
    if not branch_id:
        return False
    if is_org_wide_reporter(user) or getattr(user, 'role', None) == 'engineering_head':
        return True
    role = getattr(user, 'role', None)
    if role in ('branch_manager', 'accountant', 'auditor'):
        if getattr(user, 'branch_id', None):
            return str(user.branch_id) == str(branch_id)
        if getattr(user, 'district_id', None):
            from loans.models import Branch
            return Branch.objects.filter(pk=branch_id, district_id=user.district_id).exists()
        return False
    if role == 'district_manager' and getattr(user, 'district_id', None):
        from loans.models import Branch
        return Branch.objects.filter(pk=branch_id, district_id=user.district_id).exists()
    return False


def _district_in_user_scope(user, district_id: str) -> bool:
    if not district_id:
        return False
    if is_org_wide_reporter(user) or getattr(user, 'role', None) == 'engineering_head':
        return True
    role = getattr(user, 'role', None)
    if role == 'district_manager':
        return str(getattr(user, 'district_id', '') or '') == str(district_id)
    if role in ('branch_manager', 'accountant', 'auditor'):
        branch = getattr(user, 'branch', None)
        if branch and branch.district_id:
            return str(branch.district_id) == str(district_id)
        return str(getattr(user, 'district_id', '') or '') == str(district_id)
    return False


def apply_report_filters(qs: QuerySet, params: Dict[str, Any], user=None) -> QuerySet:
    status = (params.get('status') or '').strip()
    committee_status = (params.get('committee_status') or '').strip()
    disbursement_status = (params.get('disbursement_status') or '').strip()
    branch_id = (params.get('branch_id') or '').strip()
    district_id = (params.get('district_id') or '').strip()
    date_from = (params.get('date_from') or '').strip()
    date_to = (params.get('date_to') or '').strip()

    if status:
        qs = qs.filter(status__iexact=status)
    if committee_status:
        qs = qs.filter(committee_status=committee_status)
    if disbursement_status:
        qs = qs.filter(disbursement_status=disbursement_status)

    # Narrow only — never widen past reporting_base_queryset.
    if branch_id and (not user or _branch_in_user_scope(user, branch_id)):
        qs = qs.filter(branch_id=branch_id)

    if district_id and (not user or _district_in_user_scope(user, district_id)):
        # LO / BM with fixed branch: district filter is redundant; only apply when valid.
        role = getattr(user, 'role', None) if user else None
        if not user or is_org_wide_reporter(user) or role in ('district_manager', 'engineering_head'):
            qs = qs.filter(branch__district_id=district_id)

    if date_from:
        qs = qs.filter(date_requested__date__gte=date_from)
    if date_to:
        qs = qs.filter(date_requested__date__lte=date_to)
    return qs


def filtered_reporting_queryset(user, params: Dict[str, Any]) -> QuerySet:
    return apply_report_filters(reporting_base_queryset(user), params, user=user)


def _cell(v):
    if v is None:
        return ''
    if isinstance(v, Decimal):
        return float(v)
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d %H:%M') if hasattr(v, 'hour') else v.strftime('%Y-%m-%d')
    return v


def _autosize(ws, max_width=40):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col:
            if cell.value is not None:
                width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _write_header(ws, headers: List[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True)


def build_pipeline_workbook(qs: QuerySet) -> BytesIO:
    """Multi-sheet Excel: Pipeline + Branch summary."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pipeline'
    headers = [
        'Loan ID', 'Applicant', 'Phone', 'Branch', 'District', 'Category',
        'Amount requested', 'Status', 'Cooperative approval', 'Finance disbursement',
        'Committee status', 'Recommended amount', 'Final amount',
        'Credit score', 'Score band', 'Recommendation',
        'Disbursement status', 'Officer', 'Assigned engineer', 'Engineering status',
        'Collateral submitted', 'Sent to engineering',
        'Date requested', 'Date reviewed', 'Appraisal finished',
        'Committee decided', 'Disbursed at',
    ]
    _write_header(ws, headers)

    loans = list(qs[:5000])
    for lr in loans:
        appraisal = getattr(lr, 'appraisal', None)
        eng = getattr(lr, 'assigned_engineer', None)
        ws.append([
            _cell(lr.loan_request_id),
            _cell(lr.applicant_name),
            _cell(lr.phone_number),
            _cell(lr.branch.name if lr.branch_id else ''),
            _cell(lr.branch.district.name if lr.branch_id and lr.branch.district_id else ''),
            _cell(lr.category.name if lr.category_id else ''),
            _cell(lr.amount_requested),
            _cell(lr.status),
            'Yes' if lr.operation_manager_approval else 'No',
            'Yes' if getattr(lr, 'finance_disbursement_approval', False) else 'No',
            _cell(lr.get_committee_status_display() if lr.committee_status else ''),
            _cell(appraisal.amount_approved if appraisal else None),
            _cell(lr.committee_final_amount),
            _cell(appraisal.credit_score_total if appraisal else None),
            _cell(appraisal.credit_score_band if appraisal else None),
            _cell(appraisal.get_recommendation_display() if appraisal and appraisal.recommendation else ''),
            _cell(lr.get_disbursement_status_display() if lr.disbursement_status else ''),
            _cell(
                (lr.assigned_loan_officer.get_full_name() or lr.assigned_loan_officer.username)
                if lr.assigned_loan_officer_id else ''
            ),
            _cell((eng.get_full_name() or eng.username) if eng else ''),
            _cell(
                lr.get_collateral_engineering_status_display()
                if getattr(lr, 'collateral_engineering_status', None) else ''
            ),
            _cell(lr.collateral_submitted_at),
            _cell(lr.sent_to_engineering_at),
            _cell(lr.date_requested),
            _cell(lr.date_reviewed),
            _cell(lr.appraisal_completed_at),
            _cell(lr.committee_decided_at),
            _cell(lr.disbursed_at),
        ])
    _autosize(ws)

    # Branch summary sheet
    summary = wb.create_sheet('By branch')
    _write_header(summary, [
        'Branch', 'District', 'Loans', 'Requested amount',
        'Queue approved', 'Committee approved', 'Disbursed',
    ])
    branch_rows = (
        qs.values('branch__name', 'branch__district__name')
        .annotate(
            total=Count('id'),
            amount=Sum('amount_requested'),
            queue_ok=Count('id', filter=Q(queue_approved=True)),
            committee_ok=Count('id', filter=Q(committee_status='committee_approved')),
            disbursed=Count('id', filter=Q(disbursement_status='disbursed')),
        )
        .order_by('branch__district__name', 'branch__name')
    )
    for row in branch_rows:
        summary.append([
            row['branch__name'] or 'Unassigned',
            row['branch__district__name'] or '',
            row['total'],
            float(row['amount'] or 0),
            row['queue_ok'],
            row['committee_ok'],
            row['disbursed'],
        ])
    _autosize(summary)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def build_appraisal_pack_workbook(loan_request) -> BytesIO:
    """Per-loan Excel pack: summary, scorecard, conditions, schedule."""
    from loans.appraisal_pack import build_committee_appraisal_pack
    from loans.disbursement import final_annual_rate_pct, final_loan_amount, final_term_months

    pack = build_committee_appraisal_pack(loan_request)
    appraisal = pack.get('appraisal')
    basic = pack.get('basic_info')
    scorecard = pack.get('scorecard') or {}

    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    _write_header(summary, ['Field', 'Value'])
    rows = [
        ('Loan ID', loan_request.loan_request_id),
        ('Applicant', loan_request.applicant_name),
        ('Branch', loan_request.branch.name if loan_request.branch_id else ''),
        ('Amount requested', loan_request.amount_requested),
        ('Status', loan_request.status),
        ('Committee status', loan_request.get_committee_status_display() if loan_request.committee_status else ''),
        ('Final amount', loan_request.committee_final_amount or final_loan_amount(loan_request, appraisal)),
        ('Recommended amount', appraisal.amount_approved if appraisal else ''),
        ('Term (months)', appraisal.term_approved_months if appraisal else final_term_months(loan_request, appraisal, basic)),
        ('Rate %', appraisal.rate_approved if appraisal else final_annual_rate_pct(loan_request, appraisal, basic)),
        ('Recommendation', appraisal.get_recommendation_display() if appraisal and appraisal.recommendation else ''),
        ('Credit score', scorecard.get('total', '')),
        ('Score band', scorecard.get('band_label', '')),
        ('DSCR', getattr(appraisal, 'dscr', None) if appraisal else ''),
        ('Disbursement status', loan_request.get_disbursement_status_display() if loan_request.disbursement_status else ''),
    ]
    for label, val in rows:
        summary.append([label, _cell(val)])
    _autosize(summary)

    if scorecard.get('pillars'):
        sc = wb.create_sheet('Scorecard')
        _write_header(sc, ['Pillar', 'Earned', 'Max', 'Remark'])
        for p in scorecard['pillars']:
            sc.append([p.get('label'), p.get('earned'), p.get('max'), p.get('remark')])
        _autosize(sc)

    conditions = pack.get('conditions') or []
    if conditions:
        cd = wb.create_sheet('Conditions')
        _write_header(cd, ['Type', 'Description', 'Responsible', 'Before disbursement', 'Fulfilled'])
        for c in conditions:
            cd.append([
                c.get_condition_type_display() if hasattr(c, 'get_condition_type_display') else getattr(c, 'condition_type', ''),
                getattr(c, 'description', ''),
                getattr(c, 'responsible_party', '') or '',
                'Yes' if getattr(c, 'required_before_disbursement', False) else 'No',
                'Yes' if getattr(c, 'fulfilled', False) else 'No',
            ])
        _autosize(cd)

    amort = pack.get('amortization') or []
    if amort:
        am = wb.create_sheet('Schedule')
        _write_header(am, ['#', 'Date', 'Payment', 'Principal', 'Interest', 'Balance'])
        for row in amort:
            am.append([
                row.period_number,
                _cell(row.payment_date),
                _cell(row.payment_amount),
                _cell(row.principal),
                _cell(row.interest),
                _cell(row.balance_after),
            ])
        _autosize(am)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def excel_response(bio: BytesIO, filename: str) -> HttpResponse:
    response = HttpResponse(
        bio.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def pipeline_export_filename(params: Optional[Dict[str, Any]] = None) -> str:
    stamp = timezone.now().strftime('%Y%m%d_%H%M')
    status = (params or {}).get('status') or 'all'
    safe = ''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in str(status))[:24]
    return f'loan_pipeline_{safe}_{stamp}.xlsx'


def branch_dashboard_stats(qs: QuerySet) -> Dict[str, Any]:
    from django.db.models import Avg

    total = qs.count()
    approved = qs.filter(status__iexact='Approved').count()
    pending = qs.filter(status__iexact='Pending').count()
    rejected = qs.filter(status__iexact='Rejected').count()
    committee_pending = qs.filter(committee_status='pending_committee').count()
    committee_approved = qs.filter(committee_status='committee_approved').count()
    awaiting_disbursement = qs.filter(
        committee_status='committee_approved',
    ).exclude(disbursement_status='disbursed').count()
    disbursed = qs.filter(disbursement_status='disbursed').count()
    amount = qs.aggregate(total=Sum('amount_requested'))['total'] or Decimal('0')
    final_amt = qs.filter(committee_final_amount__isnull=False).aggregate(
        total=Sum('committee_final_amount')
    )['total'] or Decimal('0')

    scored = qs.filter(appraisal__credit_score_total__isnull=False)
    scored_n = scored.count()
    avg_score = scored.aggregate(v=Avg('appraisal__credit_score_total'))['v']
    weak_n = qs.filter(appraisal__credit_score_band__in=('weak', 'unacceptable')).count()

    by_branch = list(
        qs.values('branch_id', 'branch__name', 'branch__district__name')
        .annotate(
            total=Count('id'),
            amount=Sum('amount_requested'),
            committee_ok=Count('id', filter=Q(committee_status='committee_approved')),
            disbursed=Count('id', filter=Q(disbursement_status='disbursed')),
            pending_committee=Count('id', filter=Q(committee_status='pending_committee')),
            avg_score=Avg('appraisal__credit_score_total'),
            weak_band=Count(
                'id',
                filter=Q(appraisal__credit_score_band__in=('weak', 'unacceptable')),
            ),
        )
        .order_by('-total')[:25]
    )
    return {
        'total': total,
        'approved': approved,
        'pending': pending,
        'rejected': rejected,
        'committee_pending': committee_pending,
        'committee_approved': committee_approved,
        'awaiting_disbursement': awaiting_disbursement,
        'disbursed': disbursed,
        'amount_requested': amount,
        'final_approved_amount': final_amt,
        'approval_rate': round((approved / total) * 100, 1) if total else 0,
        'avg_credit_score': round(float(avg_score), 1) if avg_score is not None else None,
        'scored_loans': scored_n,
        'weak_band_count': weak_n,
        'weak_band_share': round((weak_n / scored_n) * 100, 1) if scored_n else 0,
        'by_branch': by_branch,
    }


def engineering_workload_stats(qs: QuerySet) -> Dict[str, Any]:
    """Collateral / engineering QA workload KPIs for eng roles (and others viewing eng columns)."""
    from loans.models import LoanRequest

    total = qs.count()
    pending = qs.filter(
        collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING,
    ).count()
    returned = qs.filter(
        collateral_engineering_status=LoanRequest.ENG_COLLATERAL_RETURNED,
    ).count()
    approved = qs.filter(
        collateral_engineering_status=LoanRequest.ENG_COLLATERAL_APPROVED,
    ).count()
    submitted = qs.filter(collateral_submitted_at__isnull=False).count()
    sent = qs.filter(sent_to_engineering_at__isnull=False).count()
    unassigned = qs.filter(
        assigned_engineer__isnull=True,
    ).filter(
        Q(sent_to_engineering_at__isnull=False)
        | Q(collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING)
    ).count()

    by_engineer = list(
        qs.exclude(assigned_engineer__isnull=True)
        .values(
            'assigned_engineer_id',
            'assigned_engineer__username',
            'assigned_engineer__first_name',
            'assigned_engineer__last_name',
        )
        .annotate(
            total=Count('id'),
            pending=Count(
                'id',
                filter=Q(collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING),
            ),
            returned=Count(
                'id',
                filter=Q(collateral_engineering_status=LoanRequest.ENG_COLLATERAL_RETURNED),
            ),
            approved=Count(
                'id',
                filter=Q(collateral_engineering_status=LoanRequest.ENG_COLLATERAL_APPROVED),
            ),
        )
        .order_by('-total')[:25]
    )
    for row in by_engineer:
        first = (row.get('assigned_engineer__first_name') or '').strip()
        last = (row.get('assigned_engineer__last_name') or '').strip()
        name = f'{first} {last}'.strip()
        row['name'] = name or row.get('assigned_engineer__username') or '—'

    return {
        'total': total,
        'pending_review': pending,
        'returned': returned,
        'approved': approved,
        'collateral_submitted': submitted,
        'sent_to_engineering': sent,
        'unassigned': unassigned,
        'by_engineer': by_engineer,
    }


def filter_choices_for_user(user, district_id=None) -> Dict[str, Any]:
    from loans.models import Branch, District, LoanRequest

    role = getattr(user, 'role', None)
    districts = District.objects.none()
    branches = Branch.objects.none()
    if is_org_wide_reporter(user) or role == 'engineering_head':
        districts = District.objects.order_by('name')
        if district_id:
            branches = Branch.objects.filter(district_id=district_id).order_by('name')
        else:
            branches = Branch.objects.none()
    elif role == 'district_manager' and getattr(user, 'district_id', None):
        districts = District.objects.filter(pk=user.district_id)
        branches = Branch.objects.filter(district_id=user.district_id).select_related('district').order_by('name')
    elif role == 'branch_manager':
        if getattr(user, 'branch_id', None):
            branches = Branch.objects.filter(pk=user.branch_id).select_related('district')
            if user.branch and user.branch.district_id:
                districts = District.objects.filter(pk=user.branch.district_id)
        elif getattr(user, 'district_id', None):
            districts = District.objects.filter(pk=user.district_id)
            branches = Branch.objects.filter(district_id=user.district_id).select_related('district').order_by('name')

    # Geo pickers only when the user can meaningfully narrow a wider scope.
    can_pick_geo = is_org_wide_reporter(user) or role in ('district_manager', 'engineering_head') or (
        role == 'branch_manager' and not getattr(user, 'branch_id', None) and getattr(user, 'district_id', None)
    )

    return {
        'districts': districts,
        'branches': branches,
        'status_choices': [
            ('', 'All statuses'),
            ('Approved', 'Approved'),
            ('Pending', 'Pending'),
            ('Rejected', 'Rejected'),
        ],
        'committee_status_choices': [('', 'All committee')] + list(LoanRequest.COMMITTEE_STATUS_CHOICES),
        'disbursement_status_choices': [('', 'All disbursement')] + [
            c for c in LoanRequest.DISBURSE_STATUS_CHOICES if c[0]
        ],
        'can_pick_geo': can_pick_geo,
        'scope_label': report_scope_label(user),
        'scope_kind': report_scope_kind(user),
        'show_engineering_workload': is_engineering_reporter(user) or is_org_wide_reporter(user),
    }
