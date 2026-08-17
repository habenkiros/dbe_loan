"""Staff-facing market actor registry + price observation capture."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from partners.forms import MarketActorForm, MarketObservationForm
from partners.market_bands import recompute_all_bands
from partners.models import MarketActor, MarketObservation, MarketPriceBand
from partners.permissions import (
    resolve_write_branch,
    scope_actors_qs,
    scope_observations_qs,
    user_can_use_market,
    user_can_write_market,
    user_may_access_actor,
)


def _forbid_if_no_access(user):
    if not user_can_use_market(user):
        return HttpResponseForbidden('Market data module is not available for your role.')
    return None


def _forbid_if_no_write(user):
    if not user_can_write_market(user):
        return HttpResponseForbidden('You cannot record market data with your role.')
    return None


@login_required
def dashboard(request):
    deny = _forbid_if_no_access(request.user)
    if deny:
        return deny

    actors = scope_actors_qs(request.user, MarketActor.objects.all())
    obs = scope_observations_qs(
        request.user,
        MarketObservation.objects.filter(status=MarketObservation.STATUS_ACTIVE),
    )
    today = timezone.localdate()
    week_ago = today - timedelta(days=7)

    week_obs = obs.filter(observed_at__gte=week_ago)
    stats = {
        'actors_active': actors.filter(is_active=True).count(),
        'actors_trusted': actors.filter(trust_status=MarketActor.TRUST_TRUSTED).count(),
        'obs_week': week_obs.count(),
        'obs_today': obs.filter(observed_at=today).count(),
        'obs_total': obs.count(),
    }
    sum_price = week_obs.aggregate(s=Sum('unit_price_etb'))['s']
    stats['sample_prices_logged'] = week_obs.count()

    recent = (
        obs.select_related('market_actor', 'city', 'recorded_by')
        .order_by('-observed_at', '-id')[:12]
    )
    by_class = (
        week_obs.values('asset_class')
        .annotate(n=Count('id'))
        .order_by('-n')
    )

    return render(request, 'partners/dashboard.html', {
        'stats': stats,
        'recent': recent,
        'by_class': by_class,
        'can_write': user_can_write_market(request.user),
        'sum_price_note': sum_price,
    })


@login_required
def actor_list(request):
    deny = _forbid_if_no_access(request.user)
    if deny:
        return deny

    qs = scope_actors_qs(
        request.user,
        MarketActor.objects.select_related('branch', 'primary_city'),
    )
    q = (request.GET.get('q') or '').strip()
    kind = (request.GET.get('kind') or '').strip()
    trust = (request.GET.get('trust') or '').strip()
    active = request.GET.get('active', '1')

    if q:
        qs = qs.filter(
            Q(name__icontains=q)
            | Q(phone_number__icontains=q)
            | Q(product_lines__icontains=q)
            | Q(contact_person__icontains=q)
        )
    if kind:
        qs = qs.filter(actor_kind=kind)
    if trust:
        qs = qs.filter(trust_status=trust)
    if active == '1':
        qs = qs.filter(is_active=True)
    elif active == '0':
        qs = qs.filter(is_active=False)

    if request.GET.get('export') == 'xlsx':
        return _export_actors_xlsx(qs)

    page = Paginator(qs.order_by('name'), 10).get_page(request.GET.get('page'))
    return render(request, 'partners/actor_list.html', {
        'page_obj': page,
        'q': q,
        'kind': kind,
        'trust': trust,
        'active': active,
        'kind_choices': MarketActor.KIND_CHOICES,
        'trust_choices': MarketActor.TRUST_CHOICES,
        'can_write': user_can_write_market(request.user),
    })


@login_required
def actor_create(request):
    deny = _forbid_if_no_write(request.user) or _forbid_if_no_access(request.user)
    if deny:
        return deny

    branch = resolve_write_branch(request.user)
    require_branch = branch is None

    if request.method == 'POST':
        form = MarketActorForm(request.POST, require_branch=require_branch)
        if form.is_valid():
            obj = form.save(commit=False)
            if branch:
                obj.branch = branch
            elif form.cleaned_data.get('branch'):
                obj.branch = form.cleaned_data['branch']
            if not obj.branch_id:
                messages.error(request, 'Branch is required.')
                return render(request, 'partners/actor_form.html', {
                    'form': form, 'is_edit': False, 'branch': branch,
                })
            obj.created_by = request.user
            obj.save()
            messages.success(request, f'Registered market actor: {obj.name}.')
            return redirect('partners:actor_detail', pk=obj.pk)
    else:
        form = MarketActorForm(
            initial={
                'trust_status': MarketActor.TRUST_PENDING,
                'product_focus': MarketActor.FOCUS_BUILDING,
            },
            require_branch=require_branch,
        )
    return render(request, 'partners/actor_form.html', {
        'form': form,
        'is_edit': False,
        'branch': branch,
    })


@login_required
def actor_edit(request, pk):
    deny = _forbid_if_no_write(request.user) or _forbid_if_no_access(request.user)
    if deny:
        return deny

    qs = scope_actors_qs(request.user, MarketActor.objects.all())
    actor = get_object_or_404(qs, pk=pk)

    if request.method == 'POST':
        form = MarketActorForm(request.POST, instance=actor)
        if form.is_valid():
            form.save()
            messages.success(request, 'Market actor updated.')
            return redirect('partners:actor_detail', pk=actor.pk)
    else:
        form = MarketActorForm(instance=actor)
    return render(request, 'partners/actor_form.html', {
        'form': form,
        'is_edit': True,
        'actor': actor,
        'branch': actor.branch,
    })


@login_required
def actor_detail(request, pk):
    deny = _forbid_if_no_access(request.user)
    if deny:
        return deny

    qs = scope_actors_qs(request.user, MarketActor.objects.select_related('branch', 'primary_city'))
    actor = get_object_or_404(qs, pk=pk)
    observations = (
        scope_observations_qs(request.user, actor.observations.all())
        .select_related('city', 'recorded_by')
        .order_by('-observed_at', '-id')[:40]
    )
    return render(request, 'partners/actor_detail.html', {
        'actor': actor,
        'observations': observations,
        'can_write': user_can_write_market(request.user),
    })


@login_required
def observation_log(request):
    """Staff quote proxy — phone-friendly price capture."""
    deny = _forbid_if_no_write(request.user) or _forbid_if_no_access(request.user)
    if deny:
        return deny

    actors_qs = scope_actors_qs(
        request.user,
        MarketActor.objects.filter(is_active=True),
    )
    initial = {}
    actor_id = request.GET.get('actor') or request.POST.get('market_actor')
    if actor_id:
        a = actors_qs.filter(pk=actor_id).first()
        if a:
            initial['market_actor'] = a.pk
            if a.primary_city_id:
                initial['city'] = a.primary_city_id
            if a.product_focus == MarketActor.FOCUS_LAND:
                initial['asset_class'] = MarketObservation.ASSET_LAND
            elif a.product_focus == MarketActor.FOCUS_VEHICLE:
                initial['asset_class'] = MarketObservation.ASSET_VEHICLE
            elif a.product_focus == MarketActor.FOCUS_MACHINERY:
                initial['asset_class'] = MarketObservation.ASSET_MACHINERY
            else:
                initial['asset_class'] = MarketObservation.ASSET_BUILDING

    if request.method == 'POST':
        form = MarketObservationForm(request.POST, actors_qs=actors_qs)
        if form.is_valid():
            obj = form.save(commit=False)
            if not user_may_access_actor(request.user, obj.market_actor):
                return HttpResponseForbidden('Actor out of scope.')
            obj.branch = obj.market_actor.branch
            obj.channel = MarketObservation.CHANNEL_STAFF
            obj.status = MarketObservation.STATUS_ACTIVE
            obj.recorded_by = request.user
            obj.save()
            recompute_all_bands()
            messages.success(
                request,
                f'Recorded {obj.item_label}: {obj.unit_price_etb} ETB from {obj.market_actor.name}.',
            )
            if request.POST.get('save_and_another'):
                return redirect(f"{request.path}?actor={obj.market_actor_id}")
            return redirect('partners:observation_list')
    else:
        form = MarketObservationForm(initial=initial or {
            'observed_at': timezone.localdate(),
            'asset_class': MarketObservation.ASSET_BUILDING,
            'condition': MarketObservation.CONDITION_NA,
        }, actors_qs=actors_qs)

    return render(request, 'partners/observation_log.html', {
        'form': form,
        'can_write': True,
    })


@login_required
def observation_list(request):
    deny = _forbid_if_no_access(request.user)
    if deny:
        return deny

    qs = scope_observations_qs(
        request.user,
        MarketObservation.objects.select_related(
            'market_actor', 'city', 'branch', 'recorded_by', 'sub_work', 'sub_sub_work',
        ),
    )
    q = (request.GET.get('q') or '').strip()
    asset = (request.GET.get('asset') or '').strip()
    status = (request.GET.get('status') or '').strip()
    date_from = (request.GET.get('from') or '').strip()
    date_to = (request.GET.get('to') or '').strip()

    if q:
        qs = qs.filter(
            Q(item_label__icontains=q)
            | Q(market_actor__name__icontains=q)
            | Q(notes__icontains=q)
        )
    if asset:
        qs = qs.filter(asset_class=asset)
    if status:
        qs = qs.filter(status=status)
    if date_from:
        qs = qs.filter(observed_at__gte=date_from)
    if date_to:
        qs = qs.filter(observed_at__lte=date_to)

    if request.GET.get('export') == 'xlsx':
        return _export_observations_xlsx(qs.order_by('-observed_at', '-id')[:5000])

    page = Paginator(qs.order_by('-observed_at', '-id'), 10).get_page(request.GET.get('page'))
    return render(request, 'partners/observation_list.html', {
        'page_obj': page,
        'q': q,
        'asset': asset,
        'status': status,
        'date_from': date_from,
        'date_to': date_to,
        'asset_choices': MarketObservation.ASSET_CHOICES,
        'status_choices': MarketObservation.STATUS_CHOICES,
        'can_write': user_can_write_market(request.user),
    })


def _export_actors_xlsx(qs):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Market actors'
    headers = [
        'Name', 'Kind', 'Phone', 'Contact', 'Focus', 'Product lines',
        'Trust', 'Active', 'Branch', 'City', 'Notes',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for a in qs.order_by('name')[:5000]:
        ws.append([
            a.name,
            a.get_actor_kind_display(),
            a.phone_number,
            a.contact_person,
            a.get_product_focus_display(),
            a.product_lines,
            a.get_trust_status_display(),
            'yes' if a.is_active else 'no',
            a.branch.name if a.branch_id else '',
            a.primary_city.name if a.primary_city_id else '',
            (a.notes or '')[:500],
        ])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = 'attachment; filename="market_actors.xlsx"'
    return resp


def _export_observations_xlsx(qs):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Observations'
    headers = [
        'Date', 'Actor', 'Branch', 'City', 'Asset class', 'Item', 'Unit',
        'Price ETB', 'Qty', 'Condition', 'Channel', 'Status', 'Recorded by', 'Notes',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for o in qs:
        ws.append([
            o.observed_at.isoformat() if o.observed_at else '',
            o.market_actor.name if o.market_actor_id else '',
            o.branch.name if o.branch_id else '',
            o.city.name if o.city_id else '',
            o.get_asset_class_display(),
            o.item_label,
            o.unit,
            float(o.unit_price_etb) if o.unit_price_etb is not None else '',
            float(o.quantity) if o.quantity is not None else '',
            o.get_condition_display(),
            o.get_channel_display(),
            o.get_status_display(),
            o.recorded_by.get_username() if o.recorded_by_id else '',
            (o.notes or '')[:500],
        ])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = 'attachment; filename="market_observations.xlsx"'
    return resp


@login_required
def band_list(request):
    deny = _forbid_if_no_access(request.user)
    if deny:
        return deny
    qs = MarketPriceBand.objects.select_related(
        'city', 'city__zone', 'sub_work', 'sub_sub_work',
    ).order_by('-sample_count', 'city__name')
    city_id = request.GET.get('city_id')
    if city_id:
        qs = qs.filter(city_id=city_id)
    page = Paginator(qs, 10).get_page(request.GET.get('page'))
    return render(request, 'partners/band_list.html', {
        'page_obj': page,
        'can_write': user_can_write_market(request.user),
    })


@login_required
def band_recompute(request):
    deny = _forbid_if_no_write(request.user) or _forbid_if_no_access(request.user)
    if deny:
        return deny
    if request.method != 'POST':
        return redirect('partners:band_list')
    n = recompute_all_bands()
    messages.success(request, f'Recomputed {n} market price band(s).')
    return redirect('partners:band_list')
