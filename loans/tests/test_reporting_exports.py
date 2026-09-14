"""Reporting / Excel export / branch dashboard smoke tests."""

from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from loans.models import (
    Branch,
    CollateralType,
    District,
    LoanAppraisal,
    LoanCategory,
    LoanRequest,
)
from loans.reporting import build_pipeline_workbook, filtered_reporting_queryset


User = get_user_model()


class ReportingExportTests(TestCase):
    def setUp(self):
        district = District.objects.create(name='Report Dist')
        self.branch = Branch.objects.create(name='Report Branch', district=district)
        collateral = CollateralType.objects.create(name='Report Coll')
        category = LoanCategory.objects.create(name='Report MSME')
        self.bm = User.objects.create_user(
            username='rep_bm', password='pass', phone_number='0911999001',
            role='branch_manager', branch=self.branch,
        )
        self.officer = User.objects.create_user(
            username='rep_lo', password='pass', phone_number='0911999002',
            role='loan_officer',
        )
        self.loan = LoanRequest.objects.create(
            loan_request_id='LR-REP-1',
            applicant_name='Report Applicant',
            phone_number='0911999003',
            amount_requested=Decimal('50000'),
            reason='WC',
            category=category,
            branch=self.branch,
            collateral=collateral,
            status='Approved',
            operation_manager_approval=True,
            finance_approval=True,
            queue_approved=True,
            assigned_loan_officer=self.officer,
            committee_status=LoanRequest.COMMITTEE_APPROVED,
            committee_final_amount=Decimal('45000'),
            disbursement_status=LoanRequest.DISBURSE_AWAITING_CONDITIONS,
        )
        LoanAppraisal.objects.create(
            loan_request=self.loan,
            created_by=self.officer,
            recommendation='approve',
            amount_approved=Decimal('45000'),
            credit_score_total=Decimal('78.5'),
            credit_score_band='acceptable',
        )

    def test_pipeline_workbook_has_sheets(self):
        qs = filtered_reporting_queryset(self.bm, {})
        bio = build_pipeline_workbook(qs)
        wb = load_workbook(BytesIO(bio.getvalue()))
        self.assertIn('Pipeline', wb.sheetnames)
        self.assertIn('By branch', wb.sheetnames)
        rows = list(wb['Pipeline'].iter_rows(values_only=True))
        self.assertTrue(any(r[0] == 'LR-REP-1' for r in rows[1:]))

    def test_generate_report_xlsx(self):
        client = Client()
        client.force_login(self.bm)
        resp = client.get(reverse('generate_report'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            'spreadsheetml.sheet',
            resp['Content-Type'],
        )
        wb = load_workbook(BytesIO(resp.content))
        self.assertIn('Pipeline', wb.sheetnames)

    def test_branch_dashboard_and_hub(self):
        client = Client()
        client.force_login(self.bm)
        hub = client.get(reverse('view_report_options'))
        self.assertEqual(hub.status_code, 200)
        self.assertContains(hub, 'Pipeline table')
        self.assertContains(hub, 'Export Excel')
        self.assertContains(hub, 'Portfolio dashboard')
        self.assertContains(hub, 'Scope snapshot')
        self.assertContains(hub, 'rpt-subnav')
        dash = client.get(reverse('branch_report_dashboard'))
        self.assertEqual(dash.status_code, 200)
        self.assertContains(dash, 'Portfolio reporting dashboard')
        self.assertContains(dash, 'LR-REP-1')
        self.assertContains(dash, 'Report Branch')  # scope label
        self.assertContains(dash, 'rpt-subnav')
        table = client.get(reverse('view_report'))
        self.assertEqual(table.status_code, 200)
        self.assertContains(table, 'Pipeline report')
        self.assertContains(table, 'hub-filters')

    def test_legal_officer_has_no_reports_nav(self):
        from loans.nav import nav_flags_for

        legal = User.objects.create_user(
            username='rep_legal', password='pass', phone_number='0911999011',
            role='legal_officer',
        )
        self.assertFalse(nav_flags_for(legal)['nav_show_reports'])
        client = Client()
        client.force_login(legal)
        resp = client.get(reverse('view_report_options'))
        self.assertNotEqual(resp.status_code, 200)

    def test_catalog_helper_has_core_cards(self):
        from loans.reporting import report_catalog_for

        titles = [
            card['title']
            for section in report_catalog_for(self.bm)
            for card in section['cards']
        ]
        self.assertIn('Pipeline table', titles)
        self.assertIn('Export Excel', titles)
        self.assertIn('Portfolio dashboard', titles)
        self.assertIn('Credit Intelligence', titles)

    def test_appraisal_pack_excel(self):
        client = Client()
        client.force_login(self.officer)
        resp = client.get(reverse('export_appraisal_pack_excel', args=[self.loan.pk]))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        self.assertIn('Summary', wb.sheetnames)

    def test_appraisal_pack_pdf_attachment(self):
        from loans.appraisal_pack_pdf import build_appraisal_pack_pdf, pdf_engine_name

        pdf_bytes, engine = build_appraisal_pack_pdf(self.loan)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        self.assertIn(engine, ('weasyprint', 'reportlab'))

        client = Client()
        client.force_login(self.officer)
        resp = client.get(reverse('export_appraisal_pack_pdf', args=[self.loan.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertTrue(resp.content.startswith(b'%PDF'))
        # Engine should at least be resolvable in this environment
        self.assertTrue(pdf_engine_name())


class ReportingScopeTests(TestCase):
    """BM / DM / LO / CEO see only their reporting scope."""

    def setUp(self):
        self.dist_a = District.objects.create(name='Scope Dist A')
        self.dist_b = District.objects.create(name='Scope Dist B')
        self.branch_a = Branch.objects.create(name='Branch A', district=self.dist_a)
        self.branch_b = Branch.objects.create(name='Branch B', district=self.dist_b)
        collateral = CollateralType.objects.create(name='Scope Coll')
        category = LoanCategory.objects.create(name='Scope Cat')

        self.officer_a = User.objects.create_user(
            username='scope_lo_a', password='pass', phone_number='0911888001',
            role='loan_officer', branch=self.branch_a,
        )
        self.officer_b = User.objects.create_user(
            username='scope_lo_b', password='pass', phone_number='0911888002',
            role='loan_officer', branch=self.branch_b,
        )
        self.bm_a = User.objects.create_user(
            username='scope_bm_a', password='pass', phone_number='0911888003',
            role='branch_manager', branch=self.branch_a,
        )
        self.dm_a = User.objects.create_user(
            username='scope_dm_a', password='pass', phone_number='0911888004',
            role='district_manager', district=self.dist_a,
        )
        self.ceo = User.objects.create_user(
            username='scope_ceo', password='pass', phone_number='0911888005',
            role='ceo',
        )
        self.auditor = User.objects.create_user(
            username='scope_auditor', password='pass', phone_number='0911888006',
            role='auditor',
        )

        self.loan_a = LoanRequest.objects.create(
            loan_request_id='LR-SCOPE-A',
            applicant_name='Applicant A',
            phone_number='0911888010',
            amount_requested=Decimal('100000'),
            reason='WC',
            category=category,
            branch=self.branch_a,
            collateral=collateral,
            status='Pending',
            assigned_loan_officer=self.officer_a,
        )
        self.loan_b = LoanRequest.objects.create(
            loan_request_id='LR-SCOPE-B',
            applicant_name='Applicant B',
            phone_number='0911888011',
            amount_requested=Decimal('200000'),
            reason='WC',
            category=category,
            branch=self.branch_b,
            collateral=collateral,
            status='Pending',
            assigned_loan_officer=self.officer_b,
        )

    def _ids(self, user, params=None):
        qs = filtered_reporting_queryset(user, params or {})
        return set(qs.values_list('loan_request_id', flat=True))

    def test_loan_officer_assigned_only(self):
        self.assertEqual(self._ids(self.officer_a), {'LR-SCOPE-A'})
        self.assertEqual(self._ids(self.officer_b), {'LR-SCOPE-B'})

    def test_branch_manager_own_branch_only(self):
        self.assertEqual(self._ids(self.bm_a), {'LR-SCOPE-A'})
        # Cannot widen via foreign branch_id
        self.assertEqual(self._ids(self.bm_a, {'branch_id': str(self.branch_b.pk)}), {'LR-SCOPE-A'})

    def test_district_manager_district_only(self):
        self.assertEqual(self._ids(self.dm_a), {'LR-SCOPE-A'})
        self.assertEqual(self._ids(self.dm_a, {'branch_id': str(self.branch_b.pk)}), {'LR-SCOPE-A'})

    def test_ceo_and_auditor_org_wide(self):
        self.assertEqual(self._ids(self.ceo), {'LR-SCOPE-A', 'LR-SCOPE-B'})
        self.assertEqual(self._ids(self.auditor), {'LR-SCOPE-A', 'LR-SCOPE-B'})

    def test_ceo_can_open_hub_and_export(self):
        client = Client()
        client.force_login(self.ceo)
        hub = client.get(reverse('view_report_options'))
        self.assertEqual(hub.status_code, 200)
        self.assertContains(hub, 'Organization-wide')
        self.assertContains(hub, 'Scope snapshot')
        export = client.get(reverse('generate_report'))
        self.assertEqual(export.status_code, 200)
        self.assertIn('spreadsheetml.sheet', export['Content-Type'])

    def test_bm_hub_shows_branch_scope(self):
        client = Client()
        client.force_login(self.bm_a)
        hub = client.get(reverse('view_report_options'))
        self.assertEqual(hub.status_code, 200)
        self.assertContains(hub, 'Branch A')
        self.assertNotContains(hub, 'LR-SCOPE-B')


class EngineeringReportingScopeTests(TestCase):
    def setUp(self):
        district = District.objects.create(name='Eng Dist')
        self.branch = Branch.objects.create(name='Eng Branch', district=district)
        collateral = CollateralType.objects.create(name='Eng Coll')
        category = LoanCategory.objects.create(name='Eng Cat')
        self.engineer = User.objects.create_user(
            username='eng_rep', password='pass', phone_number='0911777001',
            role='engineer',
        )
        self.other_engineer = User.objects.create_user(
            username='eng_other', password='pass', phone_number='0911777002',
            role='engineer',
        )
        self.eng_head = User.objects.create_user(
            username='eng_head_rep', password='pass', phone_number='0911777003',
            role='engineering_head',
        )
        self.loan_mine = LoanRequest.objects.create(
            loan_request_id='LR-ENG-MINE',
            applicant_name='Eng Applicant',
            phone_number='0911777010',
            amount_requested=Decimal('150000'),
            reason='WC',
            category=category,
            branch=self.branch,
            collateral=collateral,
            status='Pending',
            queue_approved=True,
            assigned_engineer=self.engineer,
            collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING,
            collateral_submitted_at=timezone.now(),
        )
        self.loan_other = LoanRequest.objects.create(
            loan_request_id='LR-ENG-OTHER',
            applicant_name='Other Eng Applicant',
            phone_number='0911777011',
            amount_requested=Decimal('160000'),
            reason='WC',
            category=category,
            branch=self.branch,
            collateral=collateral,
            status='Pending',
            queue_approved=True,
            assigned_engineer=self.other_engineer,
            collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING,
            collateral_submitted_at=timezone.now(),
        )
        self.loan_unassigned = LoanRequest.objects.create(
            loan_request_id='LR-ENG-UNASS',
            applicant_name='Unassigned Eng',
            phone_number='0911777012',
            amount_requested=Decimal('170000'),
            reason='WC',
            category=category,
            branch=self.branch,
            collateral=collateral,
            status='Pending',
            queue_approved=True,
            sent_to_engineering_at=timezone.now(),
            collateral_engineering_status=LoanRequest.ENG_COLLATERAL_PENDING,
        )

    def _ids(self, user):
        return set(filtered_reporting_queryset(user, {}).values_list('loan_request_id', flat=True))

    def test_engineer_assigned_only(self):
        self.assertEqual(self._ids(self.engineer), {'LR-ENG-MINE'})

    def test_engineering_head_sees_queue(self):
        ids = self._ids(self.eng_head)
        self.assertEqual(ids, {'LR-ENG-MINE', 'LR-ENG-OTHER', 'LR-ENG-UNASS'})

    def test_engineer_hub_and_workload(self):
        from loans.reporting import engineering_workload_stats, reporting_base_queryset

        client = Client()
        client.force_login(self.engineer)
        hub = client.get(reverse('view_report_options'))
        self.assertEqual(hub.status_code, 200)
        self.assertContains(hub, 'Assigned to you (engineering)')
        self.assertContains(hub, 'Engineering workload')
        self.assertContains(hub, 'Pending QA')

        stats = engineering_workload_stats(reporting_base_queryset(self.engineer))
        self.assertEqual(stats['total'], 1)
        self.assertEqual(stats['pending_review'], 1)

    def test_eng_head_export_and_dashboard(self):
        client = Client()
        client.force_login(self.eng_head)
        dash = client.get(reverse('branch_report_dashboard'))
        self.assertEqual(dash.status_code, 200)
        self.assertContains(dash, 'Engineering / collateral queue')
        self.assertContains(dash, 'By engineer')
        export = client.get(reverse('generate_report'))
        self.assertEqual(export.status_code, 200)
        self.assertIn('spreadsheetml.sheet', export['Content-Type'])
