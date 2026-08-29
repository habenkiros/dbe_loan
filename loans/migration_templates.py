"""Build fillable Excel templates for DECSI data migration."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from loans.models import (
    ApprovalCommitteeLevel,
    CollateralType,
    CustomUser,
    Department,
    LoanCategory,
    LoanRequest,
)

# (header, width, comment, dropdown_key or None)
Col = Tuple[str, int, str, str]

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
EXAMPLE_FILL = PatternFill('solid', fgColor='FFF2CC')
INSTR_FILL = PatternFill('solid', fgColor='D6EAF8')
THIN = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF'),
)

SHEET_SPECS: List[Dict] = [
    {
        'kind': 'regions',
        'title': '01_Regions',
        'columns': [
            ('name', 28, 'Region name. Example: Tigray', None),
        ],
        'examples': [{'name': 'Tigray'}],
    },
    {
        'kind': 'zones',
        'title': '02_Zones',
        'columns': [
            ('region', 22, 'Must match 01_Regions.name', None),
            ('name', 28, 'Geographic zone (not the operational district)', None),
        ],
        'examples': [{'region': 'Tigray', 'name': 'Mekelle Zone'}],
    },
    {
        'kind': 'cities',
        'title': '03_Cities',
        'columns': [
            ('region', 22, 'Optional if zone name is unique', None),
            ('zone', 22, 'Must match 02_Zones.name', None),
            ('name', 28, 'City / woreda — used for construction unit prices', None),
        ],
        'examples': [{'region': 'Tigray', 'zone': 'Mekelle Zone', 'name': 'Mekelle'}],
    },
    {
        'kind': 'districts',
        'title': '04_Districts',
        'columns': [
            ('name', 32, 'Operational district (queue / reporting scope)', None),
        ],
        'examples': [{'name': 'Mekelle District'}],
    },
    {
        'kind': 'branches',
        'title': '05_Branches',
        'columns': [
            ('district', 28, 'Must match 04_Districts.name (legacy column name: zone)', None),
            ('name', 36, 'Branch name (unique across DECSI)', None),
        ],
        'examples': [{'district': 'Mekelle District', 'name': 'Mekelle Main Branch'}],
    },
    {
        'kind': 'departments',
        'title': '06_Departments',
        'columns': [
            ('key', 18, 'cooperative | finance | credit | management | board', 'department_key'),
            ('name', 28, 'Display name', None),
            ('is_active', 12, 'TRUE / FALSE', 'yesno'),
            ('sort_order', 12, 'Display order (1, 2, 3…)', None),
        ],
        'examples': [
            {'key': 'cooperative', 'name': 'Branch Cooperative', 'is_active': 'TRUE', 'sort_order': 1},
            {'key': 'finance', 'name': 'Finance', 'is_active': 'TRUE', 'sort_order': 2},
            {'key': 'credit', 'name': 'Credit', 'is_active': 'TRUE', 'sort_order': 3},
            {'key': 'management', 'name': 'Management', 'is_active': 'TRUE', 'sort_order': 4},
            {'key': 'board', 'name': 'Board of Directors', 'is_active': 'TRUE', 'sort_order': 5},
        ],
    },
    {
        'kind': 'loan_categories',
        'title': '07_Loan_Categories',
        'columns': [
            ('name', 36, 'Product / loan type name', None),
            ('appraisal_mode', 18, 'msme (cashflow) or corporate', 'appraisal_mode'),
        ],
        'examples': [
            {'name': 'MSME Trade', 'appraisal_mode': 'msme'},
            {'name': 'Corporate Working Capital', 'appraisal_mode': 'corporate'},
        ],
    },
    {
        'kind': 'collateral_types',
        'title': '08_Collateral_Types',
        'columns': [
            ('name', 32, 'Shown on the loan request', None),
            ('kind', 16, 'building | land | movable | mixed', 'collateral_kind'),
        ],
        'examples': [
            {'name': 'Building / House', 'kind': 'building'},
            {'name': 'Land', 'kind': 'land'},
            {'name': 'Vehicle', 'kind': 'movable'},
            {'name': 'Building + Land', 'kind': 'mixed'},
        ],
    },
    {
        'kind': 'document_types',
        'title': '09_Document_Types',
        'columns': [
            ('name', 40, 'Document catalog name', None),
            ('order', 10, 'Display order (lower first)', None),
            ('is_required', 12, 'Fallback if a loan type has no pack', 'yesno'),
            ('allowed_extensions', 22, 'e.g. pdf,jpg,jpeg,png', None),
            ('max_file_size_mb', 16, 'Blank = bank default', None),
            ('enable_ocr_match', 16, 'TRUE to match name/phone/TIN from OCR', 'yesno'),
            ('require_officer_verification', 24, 'TRUE = officer must confirm', 'yesno'),
            ('for_appraisal_mode', 18, 'blank = both, or msme / corporate', 'appraisal_mode_or_all'),
            ('auth_notes', 40, 'Hint shown at upload', None),
        ],
        'examples': [
            {
                'name': 'National ID / Kebele ID', 'order': 10, 'is_required': 'TRUE',
                'allowed_extensions': 'pdf,jpg,jpeg,png', 'max_file_size_mb': 10,
                'enable_ocr_match': 'TRUE', 'require_officer_verification': 'FALSE',
                'for_appraisal_mode': '', 'auth_notes': 'Clear scan of both sides',
            },
        ],
    },
    {
        'kind': 'category_documents',
        'title': '10_Category_Documents',
        'columns': [
            ('category', 32, 'Must match 07_Loan_Categories.name', None),
            ('document_type', 40, 'Must match 09_Document_Types.name', None),
            ('is_required', 12, 'TRUE = required before collateral', 'yesno'),
            ('order', 10, 'Order within this loan type', None),
        ],
        'examples': [
            {'category': 'MSME Trade', 'document_type': 'National ID / Kebele ID', 'is_required': 'TRUE', 'order': 10},
        ],
    },
    {
        'kind': 'users',
        'title': '11_Users',
        'columns': [
            ('username', 22, 'Login id (unique)', None),
            ('first_name', 18, '', None),
            ('last_name', 18, '', None),
            ('email', 32, '', None),
            ('phone_number', 16, 'Required. Format column as Text in Excel', None),
            ('role', 22, 'See Lookups sheet', 'role'),
            ('district', 24, 'Required for branch/district staff', None),
            ('branch', 28, 'Leave blank for head-office roles', None),
            ('department', 18, 'HO: cooperative / finance / credit / management / board', 'department_key'),
            ('is_active', 12, 'TRUE / FALSE', 'yesno'),
            ('password', 18, 'Optional. Else use --default-password on import', None),
        ],
        'examples': [
            {
                'username': 'bm.mekele', 'first_name': 'Alem', 'last_name': 'Hailu',
                'email': 'alem.hailu@decsi.et', 'phone_number': '0914000001',
                'role': 'branch_manager', 'district': 'Mekelle District',
                'branch': 'Mekelle Main Branch', 'department': '',
                'is_active': 'TRUE', 'password': '',
            },
        ],
    },
    {
        'kind': 'committee_levels',
        'title': '12_Committee_Levels',
        'columns': [
            ('key', 18, 'Stable id: branch, district, head_office, management', None),
            ('name', 28, 'Display name', None),
            ('voter_scope', 16, 'branch | district | organization', 'voter_scope'),
            ('sequence_order', 14, '1 = first after officer submits', None),
            ('is_active', 12, 'TRUE / FALSE', 'yesno'),
            ('min_approvals_required', 20, 'Votes needed to pass this level', None),
            ('min_declines_required', 20, 'Votes needed to decline', None),
            ('tiebreaker_role', 20, 'Role that breaks a tie. Blank = default', 'role'),
            ('min_loan_amount', 18, 'Blank = no minimum', None),
            ('max_loan_amount', 18, 'Blank = no maximum', None),
        ],
        'examples': [
            {
                'key': 'branch', 'name': 'Branch Committee', 'voter_scope': 'branch',
                'sequence_order': 1, 'is_active': 'TRUE', 'min_approvals_required': 2,
                'min_declines_required': 2, 'tiebreaker_role': 'branch_manager',
                'min_loan_amount': '', 'max_loan_amount': 500000,
            },
        ],
    },
    {
        'kind': 'committee_members',
        'title': '13_Committee_Members',
        'columns': [
            ('level_key', 18, 'Must match 12_Committee_Levels.key', None),
            ('participant_type', 18, 'role = anyone with that role; user = named person', 'participant_type'),
            ('role', 22, 'Required when type is role', 'role'),
            ('username', 22, 'Required when type is user', None),
            ('label', 28, 'Optional display label', None),
            ('is_active', 12, 'TRUE / FALSE', 'yesno'),
        ],
        'examples': [
            {
                'level_key': 'branch', 'participant_type': 'role', 'role': 'branch_manager',
                'username': '', 'label': 'Branch manager', 'is_active': 'TRUE',
            },
        ],
    },
    {
        'kind': 'construction_catalog',
        'title': '14_Construction_Catalog',
        'columns': [
            ('main_work', 24, 'Level 1 — e.g. Foundation, Concrete, Wall', None),
            ('main_work_order', 16, 'Display order', None),
            ('sub_work', 24, 'Level 2 under main work', None),
            ('sub_work_order', 16, 'e.g. 1.1', None),
            ('sub_sub_work', 28, 'Level 3 item (optional)', None),
            ('sub_sub_order', 14, 'e.g. 1.1.1', None),
            ('unit_measure', 14, 'm², m³, pcs…', None),
            ('region', 18, 'Optional if city name is unique', None),
            ('zone', 18, 'Optional if city name is unique', None),
            ('city', 20, 'Woreda for this unit price', None),
            ('unit_price', 14, 'ETB per unit. Leave blank to add catalog only', None),
            ('effective_from', 16, 'Optional YYYY-MM-DD', None),
            ('effective_to', 16, 'Optional YYYY-MM-DD', None),
        ],
        'examples': [
            {
                'main_work': 'Foundation', 'main_work_order': 1, 'sub_work': 'Excavation',
                'sub_work_order': '1.1', 'sub_sub_work': 'Manual excavation',
                'sub_sub_order': '1.1.1', 'unit_measure': 'm³',
                'region': 'Tigray', 'zone': 'Mekelle Zone', 'city': 'Mekelle',
                'unit_price': 1500, 'effective_from': '', 'effective_to': '',
            },
        ],
    },
    {
        'kind': 'loan_requests',
        'title': '15_Loan_Requests',
        'columns': [
            ('loan_request_id', 20, 'Existing queue id (max 22 chars). Blank = auto HK-#########', None),
            ('applicant_name', 28, 'Required', None),
            ('phone_number', 16, 'Max 15 chars. Format as Text', None),
            ('email', 28, 'Optional', None),
            ('customer_number', 18, 'CBS / Temenos customer id', None),
            ('customer_history', 16, 'new or existing', 'customer_history'),
            ('category', 28, 'Must match 07_Loan_Categories.name', None),
            ('collateral', 24, 'Must match 08_Collateral_Types.name', None),
            ('amount_requested', 18, 'ETB', None),
            ('reason', 40, 'Purpose / remarks', None),
            ('status', 14, 'Pending | Approved | Rejected', 'loan_status'),
            ('district', 24, 'Optional if branch name is unique', None),
            ('branch', 28, 'Required. Must match 05_Branches.name', None),
            ('date_requested', 18, 'YYYY-MM-DD', None),
            ('origin_level', 14, 'branch or head_office', 'origin_level'),
            ('source_channel', 14, 'staff or online', 'source_channel'),
            ('assigned_loan_officer', 22, 'username from 11_Users', None),
            ('committee_status', 22, 'Blank, pending_committee, committee_approved, …', 'committee_status'),
            ('committee_final_amount', 20, 'Approved amount if already decided', None),
            ('disbursement_status', 22, 'Blank, ready_for_disbursement, disbursed, …', 'disbursement_status'),
            ('queue_approved', 14, 'TRUE to skip Cooperative intake queue', 'yesno'),
        ],
        'examples': [
            {
                'loan_request_id': '', 'applicant_name': 'Hagos Tesfay',
                'phone_number': '0914111222', 'email': '', 'customer_number': 'CIF001',
                'customer_history': 'existing', 'category': 'MSME Trade',
                'collateral': 'Building / House', 'amount_requested': 250000,
                'reason': 'Working capital', 'status': 'Pending',
                'district': 'Mekelle District', 'branch': 'Mekelle Main Branch',
                'date_requested': '2026-01-15', 'origin_level': 'branch',
                'source_channel': 'staff', 'assigned_loan_officer': 'lo.mekele1',
                'committee_status': '', 'committee_final_amount': '',
                'disbursement_status': '', 'queue_approved': 'FALSE',
            },
        ],
    },
]


def _dropdowns() -> Dict[str, List[str]]:
    roles = [k for k, _ in CustomUser.ROLE_CHOICES if k not in ('operation_manager', 'credit_committee')]
    return {
        'role': roles,
        'department_key': [k for k, _ in Department.KEY_CHOICES],
        'appraisal_mode': [LoanCategory.MODE_MSME, LoanCategory.MODE_CORPORATE],
        'appraisal_mode_or_all': ['', LoanCategory.MODE_MSME, LoanCategory.MODE_CORPORATE],
        'collateral_kind': [k for k, _ in CollateralType.KIND_CHOICES],
        'yesno': ['TRUE', 'FALSE'],
        'voter_scope': [k for k, _ in ApprovalCommitteeLevel.VOTER_SCOPE_CHOICES],
        'participant_type': ['role', 'user'],
        'customer_history': ['new', 'existing'],
        'loan_status': ['Pending', 'Approved', 'Rejected'],
        'origin_level': [LoanRequest.ORIGIN_BRANCH, LoanRequest.ORIGIN_HEAD_OFFICE],
        'source_channel': [LoanRequest.SOURCE_STAFF, LoanRequest.SOURCE_ONLINE],
        'committee_status': [k for k, _ in LoanRequest.COMMITTEE_STATUS_CHOICES],
        'disbursement_status': [k for k, _ in LoanRequest.DISBURSE_STATUS_CHOICES],
    }


def _style_header(ws: Worksheet, col_count: int) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = THIN
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(col_count)}1'


def _write_sheet(ws: Worksheet, spec: Dict, include_examples: bool) -> None:
    columns: Sequence[Col] = spec['columns']
    for idx, (header, width, comment, _dd) in enumerate(columns, start=1):
        cell = ws.cell(1, idx, header)
        ws.column_dimensions[get_column_letter(idx)].width = width
        if comment:
            cell.comment = Comment(comment, 'DECSI Loan Hub')
            cell.comment.width = 280
            cell.comment.height = 80
    _style_header(ws, len(columns))
    headers = [c[0] for c in columns]
    if include_examples:
        for r, example in enumerate(spec.get('examples') or [], start=2):
            for c, header in enumerate(headers, start=1):
                cell = ws.cell(r, c, example.get(header, ''))
                cell.fill = EXAMPLE_FILL
                cell.border = THIN
    # Leave empty fill rows for typing
    start = 2 + (len(spec.get('examples') or []) if include_examples else 0)
    for r in range(start, start + 8):
        for c in range(1, len(columns) + 1):
            ws.cell(r, c).border = THIN
    dropdowns = _dropdowns()
    max_row = 2000
    for idx, (_header, _w, _c, dd_key) in enumerate(columns, start=1):
        if not dd_key:
            continue
        values = dropdowns.get(dd_key) or []
        formula = '"' + ','.join(str(v) if v != '' else ' ' for v in values) + '"'
        if len(formula) > 250:
            continue
        dv = DataValidation(type='list', formula1=formula, allow_blank=True)
        dv.error = 'Pick a value from the list'
        dv.errorTitle = 'Invalid value'
        letter = get_column_letter(idx)
        dv.add(f'{letter}2:{letter}{max_row}')
        ws.add_data_validation(dv)


def _write_instructions(ws: Worksheet) -> None:
    ws.sheet_properties.tabColor = '1F4E79'
    lines = [
        ('DECSI Loan Hub — data migration pack', True),
        ('', False),
        ('Who fills this', False),
        ('DECSI IT / Credit Operations. Use data from the live queue system, CBS, or new lists.', False),
        ('', False),
        ('How to fill', False),
        ('1. Yellow rows on each sheet are examples. Overwrite them or delete them before import.', False),
        ('2. Do not rename sheet tabs or column headers (row 1). Extra columns are ignored.', False),
        ('3. Format phone and customer-number columns as Text so Excel does not drop leading zeros.', False),
        ('4. Import order is the sheet number (01 → 15), or import the whole workbook at once.', False),
        ('', False),
        ('Import (Docker)', False),
        ('docker compose exec web python manage.py import_migration_pack docs/migration_templates/DECSI_Migration_Pack.xlsx --default-password "ChangeMeNow!"', False),
        ('', False),
        ('Import one sheet', False),
        ('docker compose exec web python manage.py import_zones path/to/02_Zones.xlsx', False),
        ('docker compose exec web python manage.py import_users path/to/11_Users.xlsx --default-password "ChangeMeNow!"', False),
        ('', False),
        ('Flags', False),
        ('--dry-run     Validate without saving', False),
        ('--update      Update existing rows (matched by name / username / loan_request_id / key)', False),
        ('', False),
        ('Geography vs operations', False),
        ('Regions → Zones → Cities (woredas) are for collateral unit prices.', False),
        ('Districts → Branches are the operational queue structure (what the live system calls “zones” may be districts).', False),
        ('', False),
        ('What not to migrate in Excel', False),
        ('Passwords of existing staff (set a temporary password, then require change). MFA secrets. Document files / photos.', False),
        ('Appraisal sheets, GPS photos, committee votes — those start in the new factory after go-live, or are captured in-app.', False),
        ('Customers that exist in CBS — staff look them up by customer number; you can still put customer_number on loan rows.', False),
    ]
    ws.column_dimensions['A'].width = 140
    for i, (text, heading) in enumerate(lines, start=1):
        cell = ws.cell(i, 1, text)
        if heading:
            cell.font = Font(bold=True, size=16, color='1F4E79', name='Calibri')
        elif text and not text[0].isdigit() and text.endswith('e') is False and i < 8:
            cell.font = Font(bold=True, size=12, name='Calibri')
        else:
            cell.font = Font(name='Calibri', size=11)
        if heading or (text and not text.startswith('docker') and not text.startswith('--') and len(text) < 40 and i > 2):
            if text in {
                'Who fills this', 'How to fill', 'Import (Docker)', 'Import one sheet',
                'Flags', 'Geography vs operations', 'What not to migrate in Excel',
            }:
                cell.font = Font(bold=True, size=12, color='1F4E79', name='Calibri')
                cell.fill = INSTR_FILL


def _write_lookups(ws: Worksheet) -> None:
    ws.sheet_properties.tabColor = '548235'
    blocks = [
        ('Roles (use the key in 11_Users.role)', [(k, lab) for k, lab in CustomUser.ROLE_CHOICES if k not in ('operation_manager', 'credit_committee')]),
        ('Department keys', list(Department.KEY_CHOICES)),
        ('Appraisal mode', list(LoanCategory.APPRAISAL_MODE_CHOICES)),
        ('Collateral kind', list(CollateralType.KIND_CHOICES)),
        ('Committee voter_scope', list(ApprovalCommitteeLevel.VOTER_SCOPE_CHOICES)),
        ('Loan origin_level', list(LoanRequest.ORIGIN_CHOICES)),
        ('Source channel', list(LoanRequest.SOURCE_CHANNEL_CHOICES)),
        ('Committee status', list(LoanRequest.COMMITTEE_STATUS_CHOICES)),
        ('Disbursement status', list(LoanRequest.DISBURSE_STATUS_CHOICES)),
    ]
    row = 1
    ws.column_dimensions['A'].width = 36
    ws.column_dimensions['B'].width = 56
    for title, pairs in blocks:
        cell = ws.cell(row, 1, title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        ws.cell(row, 2, 'Label').font = Font(bold=True, color='FFFFFF')
        ws.cell(row, 2).fill = HEADER_FILL
        row += 1
        for key, label in pairs:
            ws.cell(row, 1, key or '(blank)')
            ws.cell(row, 2, label)
            row += 1
        row += 1


def build_pack_workbook(include_examples: bool = True) -> Workbook:
    wb = Workbook()
    instr = wb.active
    instr.title = '00_Instructions'
    _write_instructions(instr)
    lookups = wb.create_sheet('Lookups')
    _write_lookups(lookups)
    for spec in SHEET_SPECS:
        ws = wb.create_sheet(spec['title'])
        _write_sheet(ws, spec, include_examples=include_examples)
    return wb


def write_individual_files(output_dir: Path, include_examples: bool = True) -> List[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written: List[Path] = []
    for spec in SHEET_SPECS:
        wb = Workbook()
        ws = wb.active
        ws.title = spec['title']
        _write_sheet(ws, spec, include_examples=include_examples)
        path = output_dir / f"{spec['title']}.xlsx"
        wb.save(path)
        written.append(path)
    return written


def write_pack(output_dir: Path, include_examples: bool = True) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / 'DECSI_Migration_Pack.xlsx'
    build_pack_workbook(include_examples=include_examples).save(path)
    return path
